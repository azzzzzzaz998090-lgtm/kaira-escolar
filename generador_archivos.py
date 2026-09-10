import os
import re
import csv
import io
import unicodedata
from google import genai
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CARPETA_ARCHIVOS = os.path.join(BASE_DIR, "archivos_telegram")
os.makedirs(CARPETA_ARCHIVOS, exist_ok=True)

API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
client = genai.Client(api_key=API_KEY) if API_KEY else None


def normalizar(texto):
    texto = str(texto).lower().strip()
    texto = unicodedata.normalize("NFD", texto)
    return "".join(c for c in texto if unicodedata.category(c) != "Mn")


def limpiar_nombre(nombre):
    nombre = normalizar(nombre)
    nombre = re.sub(r"[^a-z0-9\s_-]", "", nombre)
    nombre = re.sub(r"\s+", "_", nombre).strip("_")
    return (nombre or "archivo_kaira")[:60]


def guardar_texto(nombre, contenido):
    ruta = os.path.join(CARPETA_ARCHIVOS, nombre)
    with open(ruta, "w", encoding="utf-8") as f:
        f.write(contenido)
    return ruta


def detectar_tipo(mensaje):
    texto = normalizar(mensaje)
    if "pdf" in texto:
        return "pdf"
    if "word" in texto or "docx" in texto:
        return "word"
    if "excel" in texto or "xlsx" in texto or "hoja de calculo" in texto:
        return "excel"
    if "arduino" in texto or ".ino" in texto:
        return "arduino"
    if "python" in texto or ".py" in texto:
        return "python"
    if "javascript" in texto or ".js" in texto:
        return "javascript"
    if "html" in texto or ".html" in texto:
        return "html"
    if "c++" in texto or "cpp" in texto:
        return "cpp"
    if "java" in texto or ".java" in texto:
        return "java"
    if "imagen" in texto or "foto" in texto or "dibujo" in texto or "poster" in texto:
        return "imagen"
    if "codigo" in texto or "programa" in texto:
        return "codigo"
    return None


def es_peticion_generacion(mensaje):
    texto = normalizar(mensaje)
    verbos = ("genera", "generame", "crea", "creame", "haz", "hazme", "prepara", "preparame", "elabora", "elaborame", "realiza", "realizame", "construye", "escribe")
    tipos = ("codigo", "programa", "pdf", "word", "excel", "imagen", "foto", "dibujo", "poster", "documento", "archivo", "arduino", "python", "javascript", "html", "c++")
    return any(v in texto for v in verbos) and any(t in texto for t in tipos)


def extraer_tema(mensaje):
    texto = str(mensaje).strip()
    texto = re.sub(r"^kaira[\s,]*", "", texto, flags=re.I)
    for patron in (r"genera(?:me)?\s+", r"crea(?:me)?\s+", r"haz(?:me)?\s+", r"prepara(?:me)?\s+", r"elabora(?:me)?\s+", r"realiza(?:me)?\s+", r"construye\s+", r"escribe\s+"):
        texto = re.sub(patron, "", texto, count=1, flags=re.I)
    texto = re.sub(r"\s+(?:y\s+)?m[aá]ndam(?:elo|ela).*?$", "", texto, flags=re.I)
    texto = re.sub(r"\s+por telegram.*?$", "", texto, flags=re.I)
    return texto.strip(" .,!?") or "archivo generado por KAIRA"


def generar_texto_gemini(prompt):
    if client is None:
        raise RuntimeError("GEMINI_API_KEY no está configurada en Render.")
    respuesta = client.interactions.create(model="gemini-3.1-flash-lite", input=prompt)
    texto = (getattr(respuesta, "output_text", "") or "").strip()
    if not texto:
        raise RuntimeError("Gemini no devolvió contenido.")
    return texto


def limpiar_codigo(codigo):
    return re.sub(r"```[a-zA-Z0-9_+.-]*\s*|```", "", codigo or "").strip()


def generar_codigo(tema, tipo):
    extensiones = {"arduino": ".ino", "python": ".py", "javascript": ".js", "html": ".html", "cpp": ".cpp", "java": ".java", "codigo": ".py"}
    lenguajes = {"arduino": "Arduino C/C++", "python": "Python", "javascript": "JavaScript", "html": "HTML", "cpp": "C++", "java": "Java", "codigo": "Python"}
    codigo = generar_texto_gemini(f"Genera código completo y funcional en {lenguajes.get(tipo, 'Python')}.\n\nSolicitud:\n{tema}\n\nEntrega únicamente código, sin Markdown. Incluye comentarios útiles y, si es Arduino, setup() y loop().")
    nombre = limpiar_nombre(tema) + extensiones.get(tipo, ".py")
    return guardar_texto(nombre, limpiar_codigo(codigo))


def generar_pdf(tema):
    contenido = generar_texto_gemini(f"Genera un documento académico en español sobre:\n{tema}\n\nIncluye título, introducción, desarrollo, conceptos importantes, ejemplos y conclusión. Usa texto claro y organizado. No uses Markdown.")
    ruta = os.path.join(CARPETA_ARCHIVOS, limpiar_nombre(tema) + ".pdf")
    estilos = getSampleStyleSheet()
    titulo = estilos["Title"]; titulo.alignment = TA_CENTER
    normal = estilos["BodyText"]; normal.leading = 15
    elementos = []
    lineas = contenido.splitlines()
    if lineas:
        elementos += [Paragraph(lineas[0], titulo), Spacer(1, 20)]
        lineas = lineas[1:]
    for linea in lineas:
        linea = linea.strip()
        if not linea:
            elementos.append(Spacer(1, 8)); continue
        safe = linea.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        estilo = estilos["Heading2"] if len(linea) < 80 and (linea.endswith(":") or linea.isupper()) else normal
        elementos.append(Paragraph(safe, estilo))
    SimpleDocTemplate(ruta, pagesize=letter, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=50).build(elementos)
    return ruta


def generar_word(tema):
    contenido = generar_texto_gemini(f"Genera el contenido completo de una investigación en español sobre:\n{tema}\n\nIncluye introducción, desarrollo, conceptos principales, características, ejemplos, aplicaciones, ventajas y desventajas cuando correspondan y conclusión. No uses Markdown.")
    ruta = os.path.join(CARPETA_ARCHIVOS, limpiar_nombre(tema) + ".docx")
    doc = Document()
    lineas = contenido.splitlines()
    if lineas:
        doc.add_heading(lineas[0], 0); lineas = lineas[1:]
    for linea in lineas:
        linea = linea.strip()
        if not linea: continue
        if len(linea) < 80 and (linea.endswith(":") or linea.isupper()):
            doc.add_heading(linea, level=1)
        else:
            doc.add_paragraph(linea)
    doc.save(ruta)
    return ruta


def generar_excel(tema):
    contenido = generar_texto_gemini(f"Crea una tabla para Excel sobre:\n{tema}\n\nDevuelve únicamente CSV, primera fila con encabezados, 8 a 15 filas, sin Markdown.")
    filas = list(csv.reader(io.StringIO(limpiar_codigo(contenido))))
    if not filas: filas = [["Información"], [contenido]]
    ruta = os.path.join(CARPETA_ARCHIVOS, limpiar_nombre(tema) + ".xlsx")
    libro = Workbook(); hoja = libro.active; hoja.title = "Datos"
    for fila in filas: hoja.append(fila)
    for celda in hoja[1]: celda.font = Font(bold=True); celda.alignment = Alignment(horizontal="center")
    for columna in hoja.columns:
        letra = columna[0].column_letter
        maximo = max((len(str(c.value)) for c in columna if c.value is not None), default=10)
        hoja.column_dimensions[letra].width = min(maximo + 3, 50)
    libro.save(ruta)
    return ruta


def procesar_peticion_generacion(mensaje):
    if not es_peticion_generacion(mensaje):
        return None
    tipo = detectar_tipo(mensaje)
    if not tipo:
        return None
    tema = extraer_tema(mensaje)
    try:
        if tipo in {"arduino", "python", "javascript", "html", "cpp", "java", "codigo"}:
            ruta = generar_codigo(tema, tipo)
        elif tipo == "pdf":
            ruta = generar_pdf(tema)
        elif tipo == "word":
            ruta = generar_word(tema)
        elif tipo == "excel":
            ruta = generar_excel(tema)
        else:
            raise RuntimeError("La generación de imágenes requiere un motor de imágenes configurado; no se intenta usar FastSD local en Render.")
        return {"ruta": ruta, "nombre": os.path.basename(ruta), "tipo": tipo, "tema": tema}
    except Exception as error:
        print("❌ ERROR GENERANDO ARCHIVO:", type(error).__name__, error)
        return {"error": str(error)}
